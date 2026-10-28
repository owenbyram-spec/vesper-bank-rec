"""Exact-session exports and bounded, checksummed workspace backups."""

import csv
import html
import io
import json
import re
from .models import DomainError, digest, money, day, label, uid, now
from .reconciliation import SCHEMA_VERSION, summary
from .matching import validate_pair

NUMERIC = re.compile(r"^-?\d+(?:\.\d+)?$")


def safe(value):
    text = "" if value is None else str(value)
    if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) and not NUMERIC.fullmatch(text):
        return "'" + text
    return text


def export_rows(workspace, session):
    result = summary(session)
    rows = [["section", "key / date", "value / description", "amount", "source / row ID"]]
    rows += [
        ["context", "client", workspace["name"], "", workspace["id"]],
        ["context", "synthetic", str(workspace["synthetic"]), "", ""],
        ["context", "account", session["account_label"], "", session["account_id"]],
        ["context", "session", session["id"], "", ""],
        ["context", "period", f"{session['period_start']} to {session['period_end']}", "", ""],
        ["context", "currency / status", f"{session['currency']} / {session['status']}", "", ""],
        ["context", "input version", str(session["input_version"]), "", result["evidence_id"]],
    ]
    for key in (
        "bank_closing",
        "deposits_in_transit",
        "outstanding_withdrawals",
        "bank_corrections",
        "adjusted_bank",
        "book_closing",
        "pending_book_adjustments",
        "adjusted_books",
        "difference",
    ):
        rows.append(["tie_out", key, "", result[key], ""])
    anchor = session.get("anchor") or {}
    rows.append(
        [
            "opening_anchor",
            anchor.get("kind", "missing"),
            anchor.get("evidence", ""),
            "",
            anchor.get("as_of", ""),
        ]
    )
    for side in ("bank", "ledger"):
        source = session.get(side)
        if not source:
            continue
        rows.extend(
            [
                ["source", side, source["source"], "", source["source_hash"]],
                ["control", f"{side} opening", "", source["opening"], source["observed_at"]],
                ["control", f"{side} activity", "", source["activity"], ""],
                ["control", f"{side} closing", "", source["closing"], ""],
            ]
        )
        if source.get("provider_provenance"):
            rows.append(
                [
                    "provider_provenance",
                    side,
                    json.dumps(source["provider_provenance"], sort_keys=True),
                    "",
                    "",
                ]
            )
        for row in source["rows"]:
            rows.append(
                [
                    side,
                    row["date"],
                    row["description"],
                    row["amount"],
                    f"{row['id']} / ordinal {row['ordinal']}",
                ]
            )
    for row in result["outstanding"]:
        rows.append(["outstanding", row["date"], row["description"], row["amount"], row["id"]])
        rows.append(
            ["outstanding_evidence", row["id"], row.get("evidence", ""), "", row.get("carried_from", "")]
        )
    for row in result["pending_adjustments"]:
        rows.append(
            ["pending_external_posting", row["date"], row["description"], row["amount"], row["evidence"]]
        )
    for row in result["corrections"]:
        rows.append(["bank_correction", row["date"], row["description"], row["adjustment"], row["evidence"]])
    for match in session["matches"]:
        rows.append(["human_match", match["bank_id"], match["ledger_id"], "", match["accepted_at"]])
    for blocker in result["blockers"]:
        rows.append(["completion_blocker", "", blocker, "", ""])
    for entry in session["history"]:
        rows.append(
            ["history", entry["at"], entry["action"], "", json.dumps(entry["detail"], sort_keys=True)]
        )
    return rows


def csv_bytes(workspace, session):
    stream = io.StringIO(newline="")
    csv.writer(stream).writerows([[safe(cell) for cell in row] for row in export_rows(workspace, session)])
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def xlsx_bytes(workspace, session):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reconciliation"
    for row in export_rows(workspace, session):
        sheet.append([safe(cell) for cell in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="172A3B")
    for column, width in {"A": 27, "B": 40, "C": 65, "D": 22, "E": 65}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def html_report(workspace, session):
    def escape(value):
        return html.escape(str(value))

    result = summary(session)
    body = "".join(
        "<tr>"
        + "".join(
            f'<td class="{"numeric" if NUMERIC.fullmatch(str(cell)) else "text"}">{escape(cell)}</td>'
            for cell in row
        )
        + "</tr>"
        for row in export_rows(workspace, session)[1:]
    )
    return f"""<!doctype html><html lang="en"><meta charset="utf-8"><title>Vesper reconciliation report</title>
<style>body{{font:14px system-ui;color:#172a3b;max-width:1200px;margin:40px auto;padding:20px}}h1{{font-family:Georgia}}table{{width:100%;border-collapse:collapse;table-layout:fixed}}th:nth-child(1){{width:12%}}th:nth-child(2){{width:18%}}th:nth-child(3){{width:31%}}th:nth-child(4){{width:16%}}th:nth-child(5){{width:23%}}td.numeric{{white-space:nowrap}}td,th{{border-bottom:1px solid #ddd;text-align:left;padding:7px;overflow-wrap:anywhere}}thead{{display:table-header-group}}@media print{{body{{margin:0;font-size:10px}}tr{{break-inside:avoid}}}}</style>
<h1>Vesper &amp; Co · Bank reconciliation</h1><p>{escape(workspace["name"])} · {escape(session["account_label"])} · {escape(session["period_start"])}–{escape(session["period_end"])}</p>
<p>{"SYNTHETIC EXAMPLE — fictional data." if workspace["synthetic"] else "Local workspace — entered names are not verified identities."} Status: {escape(session["status"])}. Input version {session["input_version"]}.</p>
<p>Adjusted bank <strong>{result["adjusted_bank"]}</strong> · Adjusted books <strong>{result["adjusted_books"]}</strong> · Difference <strong>{result["difference"]} {escape(session["currency"])}</strong></p>
<p>Use your browser’s Print → Save as PDF. A draft report is not a completed reconciliation.</p>
<table><thead><tr><th>Section</th><th>Key / date</th><th>Value / description</th><th>Amount</th><th>Source / row ID</th></tr></thead><tbody>{body}</tbody></table></html>"""


def backup(workspace):
    payload = {
        "schema_version": SCHEMA_VERSION,
        "product": "vesper-bank-rec",
        "exported_at": now(),
        "workspace": workspace,
    }
    return {**payload, "checksum": digest(payload)}


def validate_backup(payload):
    if not isinstance(payload, dict) or set(payload) != {
        "schema_version",
        "product",
        "exported_at",
        "workspace",
        "checksum",
    }:
        raise DomainError("This is not a supported Vesper Bank Rec backup.")
    body = {key: value for key, value in payload.items() if key != "checksum"}
    if (
        payload["product"] != "vesper-bank-rec"
        or payload["schema_version"] != SCHEMA_VERSION
        or digest(body) != payload["checksum"]
    ):
        raise DomainError("Backup format/checksum failed. No workspace was restored.")
    workspace = payload["workspace"]
    if (
        not isinstance(workspace, dict)
        or not isinstance(workspace.get("sessions"), list)
        or len(workspace["sessions"]) > 240
    ):
        raise DomainError("Invalid or oversized backup workspace.")
    label(workspace.get("name"), "Workspace name")
    if workspace.get("currency") not in ("USD", "CAD", "GBP", "EUR", "AUD", "NZD"):
        raise DomainError("Invalid backup currency.")
    ids = set()
    for session in workspace["sessions"]:
        if session["id"] in ids:
            raise DomainError("Repeated session identity in backup.")
        ids.add(session["id"])
        day(session["period_start"])
        day(session["period_end"])
        if session["currency"] != workspace["currency"]:
            raise DomainError("Backup contains incompatible currencies.")
        row_ids = set()
        for side in ("bank", "ledger"):
            source = session.get(side)
            if not source:
                continue
            if len(source["rows"]) > 10000:
                raise DomainError("Backup source exceeds row limits.")
            money(source["opening"])
            money(source["closing"])
            for row in source["rows"]:
                money(row["amount"])
                day(row["date"])
                if (
                    row["id"] in row_ids
                    or row["account_id"] != session["account_id"]
                    or row["currency"] != session["currency"]
                ):
                    raise DomainError("Invalid row scope or repeated identity in backup.")
                row_ids.add(row["id"])
                if not session["period_start"] <= row["date"] <= session["period_end"]:
                    raise DomainError("Backup row falls outside the session period.")
        for row in session["carry"]:
            money(row["amount"])
            day(row["date"])
            if row.get("account_id") != session["account_id"] or row.get("currency") != session["currency"]:
                raise DomainError("A carried item belongs to a different account or currency.")
            if row["id"] in row_ids or row["date"] >= session["period_start"]:
                raise DomainError("Invalid carried item in backup.")
            row_ids.add(row["id"])
        for row_id, classification in session["classifications"].items():
            if row_id not in row_ids or classification["kind"] not in (
                "outstanding",
                "bank_correction",
                "bank_fee",
                "interest",
                "returned_payment",
                "book_adjustment",
            ):
                raise DomainError("Invalid backup classification.")
            label(classification["evidence"], "Evidence", 2000)
        anchor = session.get("anchor")
        if anchor:
            from datetime import date, timedelta

            if anchor.get("kind") not in ("zero_outstanding", "prior_reconciled", "prior_completed"):
                raise DomainError("Unsupported opening anchor in backup.")
            label(anchor.get("evidence"), "Anchor evidence", 2000)
            if (
                anchor.get("as_of")
                != (date.fromisoformat(session["period_start"]) - timedelta(days=1)).isoformat()
            ):
                raise DomainError("Backup opening anchor has the wrong date.")
            money(anchor["bank_opening"])
            money(anchor["book_opening"])
            if anchor["kind"] == "zero_outstanding" and (
                session["carry"] or money(anchor["bank_opening"]) != money(anchor["book_opening"])
            ):
                raise DomainError(
                    "A zero-outstanding opening must have equal independent balances and no carried items."
                )
        elif session["carry"]:
            raise DomainError("Carried items require an evidenced opening anchor.")
        if session.get("status") not in ("draft", "review", "completed"):
            raise DomainError("Unsupported reconciliation status in backup.")
        matches = list(session["matches"])
        session["matches"] = []
        for match in matches:
            validate_pair(session, match["bank_id"], match["ledger_id"])
            session["matches"].append(match)
        result = summary(session)
        if session["status"] == "completed" and (
            not result["ready_to_complete"]
            or not session["completion"]
            or session["completion"]["result_hash"] != digest(result)
        ):
            raise DomainError("Completed backup period does not reproduce its saved exact result.")
    by_id = {session["id"]: session for session in workspace["sessions"]}
    for session in workspace["sessions"]:
        anchor = session.get("anchor") or {}
        if anchor.get("kind") != "prior_completed":
            continue
        previous = by_id.get(anchor.get("prior_session_id"))
        if (
            not previous
            or previous["status"] != "completed"
            or previous["account_id"] != session["account_id"]
        ):
            raise DomainError(
                "Backup prior-period anchor does not refer to a completed period in the same account."
            )
        if (
            anchor["as_of"] != previous["period_end"]
            or anchor["prior_result_hash"] != previous["completion"]["result_hash"]
            or anchor["bank_opening"] != previous["bank"]["closing"]
            or anchor["book_opening"] != previous["ledger"]["closing"]
        ):
            raise DomainError("Backup prior-period opening controls or result hash do not reproduce.")
        expected = {row["id"]: row for row in summary(previous)["outstanding"]}
        actual = {row["id"]: row for row in session["carry"]}
        if set(expected) != set(actual) or any(
            any(
                actual[key].get(field) != row.get(field)
                for field in ("amount", "date", "description", "account_id", "currency", "source_hash")
            )
            for key, row in expected.items()
        ):
            raise DomainError("Backup outstanding continuity does not reproduce the prior completed period.")
    if workspace["active_session_id"] is not None and workspace["active_session_id"] not in ids:
        raise DomainError("Backup active session does not exist.")
    workspace["id"] = uid("workspace")
    workspace["name"] = workspace["name"][:140] + " (restored)"
    return workspace
