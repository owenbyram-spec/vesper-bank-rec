"""Check the generated Excel file, including stored numeric values and identities."""

from decimal import Decimal
import io
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
import pytest

from vesper_bank_rec.imports import preview_csv
from vesper_bank_rec.reconciliation import commit_import, new_session
from vesper_bank_rec.reports import export_rows, xlsx_bytes


def exported_session(amounts, description="=1+1"):
    workspace = {
        "id": "00012345",
        "name": "@Synthetic client",
        "currency": "USD",
        "synthetic": True,
        "sessions": [],
        "events": [],
    }
    session = new_session(workspace, {
        "account_id": "00067890",
        "account_label": "+Synthetic checking",
        "period_start": "2026-09-01",
        "period_end": "2026-09-30",
    })
    csv = "date,description,amount,account,currency\n" + "".join(
        f"2026-09-12,{description},{value},00067890,USD\n" for value in amounts
    )
    closing = format(sum((Decimal(value) for value in amounts), Decimal(0)), ".2f")
    for side in ("bank", "ledger"):
        preview = preview_csv(session, side, {
            "csv": csv, "opening": "0.00", "closing": closing,
            "source": "-Synthetic independent controls",
        })
        commit_import(workspace, session, preview)
    return workspace, session


@pytest.mark.parametrize("value", [
    "0.00", "0.01", "-0.01", "0.10", "-75.00", "123456789012.34",
    "999999999999.98", "999999999999.99", "-999999999999.99",
])
def test_xlsx_money_round_trips_as_numeric_exact_cents(value):
    workspace, session = exported_session([] if value == "0.00" else [value])
    expected = export_rows(workspace, session)
    output = xlsx_bytes(workspace, session)
    workbook = load_workbook(io.BytesIO(output))
    sheet = workbook.active
    with ZipFile(io.BytesIO(output)) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    xml_cells = {cell.attrib["r"]: cell for cell in root.findall(".//s:c", ns)}

    for row_number, row in enumerate(expected, 1):
        for column_number, original in enumerate(row, 1):
            cell = sheet.cell(row_number, column_number)
            assert cell.data_type != "f"
            if row_number > 1 and column_number == 4 and original != "":
                assert cell.data_type == "n"
                assert cell.number_format == "#,##0.00"
                assert Decimal(str(cell.value)).quantize(Decimal(".01")) == Decimal(original)
                # Inspect the actual saved XML too: an intermediate float must
                # not discard cents near the maximum supported amount.
                stored = xml_cells[cell.coordinate]
                assert stored.attrib.get("t") == "n"
                assert Decimal(stored.find("s:v", ns).text) == Decimal(original)
            elif original != "":
                assert cell.data_type == "s"

    assert sheet["C2"].value == "'@Synthetic client"
    assert sheet["E2"].value == "00012345"
    assert sheet["C4"].value == "'+Synthetic checking"
    assert sheet["E4"].value == "00067890"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == sheet.dimensions


def test_xlsx_preserves_duplicate_occurrence_ids_and_neutralizes_text():
    workspace, session = exported_session(["-75.00", "-75.00"])
    workbook = load_workbook(io.BytesIO(xlsx_bytes(workspace, session)))
    sheet = workbook.active
    for side in ("bank", "ledger"):
        rows = [row for row in sheet if row[0].value == side]
        assert len(rows) == 2
        assert [row[2].value for row in rows] == ["'=1+1", "'=1+1"]
        assert [row[3].value for row in rows] == [-75, -75]
        assert [row[4].value for row in rows] == [
            f"{record['id']} / ordinal {record['ordinal']}" for record in session[side]["rows"]
        ]
        assert rows[0][4].value != rows[1][4].value
    assert all(cell.data_type != "f" for row in sheet for cell in row)
